import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def apply_yolo_header(doc):
  # 1. Title
  title_p = doc.add_paragraph()
  title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
  title_p.paragraph_format.space_after = Pt(12)
  t_run = title_p.add_run(
      "Benchmarking Variational Quantum Classifiers Against Classical Machine"
      " Learning Baselines on Dimensionality-Reduced Bounding-Box Features"
  )
  t_run.font.name = "Times New Roman"
  t_run.font.size = Pt(16)
  t_run.font.bold = True
  t_run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

  # 2. Author Name
  author_p = doc.add_paragraph()
  author_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
  author_p.paragraph_format.space_after = Pt(3)
  a_run = author_p.add_run("Sharath Chandran")
  a_run.font.name = "Times New Roman"
  a_run.font.size = Pt(12)

  # 3. Independent Affiliation
  affil_p = doc.add_paragraph()
  affil_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
  affil_p.paragraph_format.space_after = Pt(4)
  aff_run = affil_p.add_run("Independent Researcher")
  aff_run.font.name = "Times New Roman"
  aff_run.font.size = Pt(11)
  aff_run.font.italic = True

  # 4. Project / Code Repository URL (Matching YOLO pink/magenta style)
  url_p = doc.add_paragraph()
  url_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
  url_p.paragraph_format.space_after = Pt(18)
  u_run = url_p.add_run("https://github.com/sharath/qml-yolov8-benchmark")
  u_run.font.name = "Courier New"
  u_run.font.size = Pt(10.5)
  u_run.font.color.rgb = RGBColor(0xE6, 0x32, 0x82)  # YOLO pink/magenta